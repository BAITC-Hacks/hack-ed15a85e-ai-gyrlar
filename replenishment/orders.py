"""Supplier purchase orders: validation, commitments and portable exports."""
from __future__ import annotations

import copy
import csv
import io
import math
import uuid
from datetime import date, datetime, timedelta

STATUSES = {'draft': 'Черновик', 'approved': 'Размещён', 'received': 'Получен', 'cancelled': 'Отменён'}


def now():
    return datetime.now().isoformat(timespec='seconds')


def validate_quantity(value, pack, moq):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not 0 <= value <= 10**9:
        raise ValueError('Количество должно быть числом от 0 до 1 000 000 000.')
    if value and (value < moq or abs(value / pack-round(value / pack)) > 1e-7):
        raise ValueError('Количество не соответствует минимальной партии или кратности.')
    return value


def summary(orders):
    return {status: sum(o['status'] == status for o in orders) for status in STATUSES}


def reservations(orders):
    result = {}
    for order in orders:
        if order['status'] == 'draft':
            for line in order['lines']:
                result[line['item_id']] = result.get(line['item_id'], 0)+line['qty']
    return result


def apply_commitments(dataset, orders):
    """Work on a copied dataset; external order references reconcile the same supply."""
    lookup = {i['id']: i for i in dataset['items']}
    references = {ref: o for o in orders for ref in (o['id'], o['number'])}
    for item in dataset['items']:
        item['transit'] = [t for t in item.get('transit', [])
                           if str(t.get('order_id', '')) not in references
                           or references[str(t.get('order_id', ''))]['status'] not in ('received', 'cancelled')]
    for order in orders:
        if order['status'] != 'approved':
            continue
        for line in order['lines']:
            item = lookup.get(line['item_id'])
            if not item:
                continue
            represented = sum(t['qty'] for t in item.get('transit', [])
                              if str(t.get('order_id', '')) in (order['id'], order['number']))
            remaining = max(0, line['qty']-represented)
            if remaining:
                item.setdefault('transit', []).append(dict(qty=remaining, eta=order['expected_at'],
                    source='Заказ '+order['number'], order_id=order['number'], internal=True))


def new_order(supplier, rows, as_of, sequence, fingerprint, settings):
    expected = date.fromisoformat(as_of)+timedelta(days=max(r['lead_days'] for r in rows))
    lines = []
    for row in rows:
        qty = row.get('create_qty', row['qty'])
        lines.append({**{k: row[k] for k in ('code', 'article', 'name', 'unit', 'pack', 'moq', 'reason')},
                      'item_id': row['id'], 'qty': qty, 'recommended_qty': row['qty'],
                      'basis': fingerprint(row, settings), 'warnings': list(row['warnings'])})
    return dict(id=uuid.uuid4().hex, number=f'SP-{date.today():%Y%m%d}-{sequence:04}', supplier=supplier,
                status='draft', created_at=now(), updated_at=now(), expected_at=expected.isoformat(),
                as_of=as_of, note='', lines=lines, events=[dict(at=now(), action='Создан черновик')],
                warnings=list(dict.fromkeys(w for line in lines for w in line['warnings'])))


def safe_text(value):
    text = str(value or '')
    return "'"+text if text.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')) else text


def export_order(order, kind='xlsx'):
    header = ['Код товара', 'Артикул', 'Наименование', 'Количество', 'Единица', 'Кратность', 'Минимальная партия', 'Обоснование']
    values = [[safe_text(l['code']), safe_text(l['article']), safe_text(l['name']), l['qty'],
               safe_text(l['unit']), l['pack'], l['moq'], safe_text(l['reason'])] for l in order['lines']]
    if kind == 'csv':
        stream = io.StringIO(newline='')
        writer = csv.writer(stream, delimiter=';')
        writer.writerow(['Заказ', order['number'], 'Поставщик', safe_text(order['supplier'])])
        writer.writerow(['Статус', STATUSES[order['status']], 'Поступление', order['expected_at']])
        writer.writerow(header)
        writer.writerows(values)
        return ('\ufeff'+stream.getvalue()).encode('utf-8'), 'text/csv; charset=utf-8'
    if kind != 'xlsx':
        raise ValueError('Доступны форматы xlsx и csv.')
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise ValueError('Для Excel установите openpyxl из requirements.txt; CSV доступен без него.') from None
    book = Workbook()
    sheet = book.active
    sheet.title = 'Заказ поставщику'
    sheet.append(['Заказ '+order['number']])
    sheet.append(['Поставщик', safe_text(order['supplier']), 'Статус', STATUSES[order['status']]])
    sheet.append(['Поступление', date.fromisoformat(order['expected_at']), 'Дата создания', order['created_at'][:10]])
    sheet['B3'].number_format = 'dd.mm.yyyy'
    sheet.append(['Примечание', safe_text(order.get('note', ''))])
    sheet.append(header)
    for row in values:
        sheet.append(row)
    sheet.freeze_panes = 'D6'
    sheet.auto_filter.ref = f'A5:H{sheet.max_row}'
    for cell in sheet[5]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='19785E')
        cell.alignment = Alignment(wrap_text=True)
    sheet['A1'].font = Font(size=18, bold=True, color='152B37')
    for col, width in zip('ABCDEFGH', [21, 24, 48, 16, 12, 14, 19, 65]):
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=6):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        for col in (4, 6, 7):
            row[col-1].number_format = '#,##0.###'
    sheet.row_dimensions[5].height = 30
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:5'
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
