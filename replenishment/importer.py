"""Read the 12 partner workbooks without changing them. Keep source provenance."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, date, timedelta
from pathlib import Path

from openpyxl import load_workbook
from .engine import clean_transactions, number

MONTHS = {'янв': 1, 'фев': 2, 'мар': 3, 'апр': 4, 'май': 5, 'июн': 6,
          'июл': 7, 'авг': 8, 'сен': 9, 'окт': 10, 'ноя': 11, 'дек': 12}


def month_label(text):
    text = str(text).lower()
    year = re.search(r'20\d{2}', text)
    key = text[:3]
    return f'{year.group()}-{MONTHS[key]:02}' if year and key in MONTHS else None


def read_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        yield from wb.active.iter_rows(values_only=True)
    finally:
        wb.close()


def import_workbooks(root: Path):
    catalog = {}
    manifest = []
    seasonality = {}
    diagnostics = Counter()

    def item(supplier, code, name=''):
        key = f'{supplier}:{str(code).strip()}'
        if key not in catalog:
            catalog[key] = dict(id=key, supplier=supplier, code=str(code).strip(), name=str(name or '').strip(),
                                article='', category='Без категории', unit='', sales={}, stock_history={},
                                stock=None, stock_date=None, pack=None, moq=None, transit=[], sources=[], warnings=[])
        return catalog[key]

    for dirname, supplier in [('Systeme electric', 'Systeme Electric'), ('IEK', 'IEK')]:
        folder = root/dirname
        paths = sorted(folder.rglob('*.xlsx'))
        if len(paths) != 6:
            raise ValueError(f'{folder}: ожидались 6 файлов xlsx, найдено {len(paths)}')
        bytype = {}
        for p in paths:
            key = ('events' if p.name.startswith('Динамика') else 'stock' if p.name.startswith('Ежемесячные остатки')
                   else 'sales' if p.name.startswith('Ежемесячные продажи') else 'season' if p.name.startswith('Сезонность')
                   else 'moq' if p.name.startswith('MOQ') else 'transit')
            if key in bytype:
                raise ValueError(f'Два файла одного типа: {key}')
            bytype[key] = p
            manifest.append(dict(supplier=supplier, type=key, file=p.name,
                                 sha256=hashlib.sha256(p.read_bytes()).hexdigest(), bytes=p.stat().st_size))

        for kind in ['sales', 'stock']:
            p = bytype[kind]
            rows = iter(read_rows(p))
            header = next(rows)
            codeidx = header.index('Номенклатура.Код')
            nameidx = header.index('Номенклатура')
            monthcols = [(i, month_label(v)) for i, v in enumerate(header) if month_label(v)]
            unitidx = next((i for i, v in enumerate(header) if v in ['Ед.', 'Ед.изм']), None)
            seen = set()
            for lineno, row in enumerate(rows, 2):
                if not row[codeidx]:
                    continue
                code = str(row[codeidx]).strip()
                if code in seen:
                    raise ValueError(f'Повтор кода {code} в {p.name}:{lineno}')
                seen.add(code)
                obj = item(supplier, code, row[nameidx])
                obj['sources'].append(f'{p.name} · Лист_1 · строка {lineno}')
                if unitidx is not None and row[unitidx]:
                    obj['unit'] = str(row[unitidx])
                for col, month in monthcols:
                    obj['sales' if kind == 'sales' else 'stock_history'][month] = (
                        number(row[col]) if kind == 'sales' else number(row[col], None))
                if kind == 'sales' and 'Артикул' in header:
                    obj['article'] = str(row[header.index('Артикул')] or '').strip()
                if kind == 'stock':
                    last_month = max(m for _, m in monthcols)
                    obj['stock'] = obj['stock_history'][last_month]
                    obj['stock_date'] = last_month+'-01'
                    obj['stock_basis'] = 'Месячный остаток; снимок на начало месяца требует сверки'

        p = bytype['moq']
        for lineno, row in enumerate(read_rows(p), 1):
            if lineno == 1:
                header = row
                ci = header.index('Номенклатура.Код') if supplier == 'Systeme Electric' else header.index('Код 1с')
                ai = 3 if supplier == 'Systeme Electric' else 2
                ni = 1 if supplier == 'Systeme Electric' else 3
                continue
            if not row[ci]:
                continue
            obj = item(supplier, row[ci], row[ni])
            obj['article'] = str(row[ai] or '').strip()
            value = number(row[4], None)
            if supplier == 'Systeme Electric':
                obj['pack'] = value if value and value > 0 else None
            else:
                # IEK calls this field a minimum shipment, not a multiple.
                obj['moq'] = value if value and value > 0 else None
                obj['pack'] = 1
            obj['sources'].append(f'{p.name} · строка {lineno}')

        p = bytype['season']
        rows = list(read_rows(p))
        indices = [number(rows[i][11], 1) for i in range(10, 22)]
        if any(v <= 0 for v in indices):
            raise ValueError(f'{p.name}: неверная сезонность')
        seasonality[supplier] = indices

        p = bytype['transit']
        rows = list(read_rows(p))
        seen = set()
        for lineno, row in enumerate(rows[2:] if supplier == 'Systeme Electric' else rows[1:], 3 if supplier == 'Systeme Electric' else 2):
            ci, ai, ni = (2, 1, 3) if supplier == 'Systeme Electric' else (0, 1, 2)
            if not row[ci]:
                continue
            obj = item(supplier, row[ci], row[ni])
            if obj['id'] in seen:
                diagnostics['duplicate_transit_codes'] += 1
                obj['warnings'].append('Код повторяется в поставках: количества суммируются')
            seen.add(obj['id'])
            obj['article'] = str(row[ai] or obj['article']).strip()
            obj['sources'].append(f'{p.name} · строка {lineno}')
            if supplier == 'Systeme Electric':
                obj['category'] = 'Категория '+str(row[4] or 'не указана')
                obj['source_growth'] = number(row[43], None)
                obj['stock'] = number(row[51], None)
                obj['stock_date'] = '2026-09-22'
                obj['stock_basis'] = 'Свободный остаток из ведомости на 22.09.2026'
                if number(row[54]) > 0:
                    obj['transit'].append(dict(qty=number(row[54]), eta='2026-09-24', source=p.name))
                # Retain every product from the statement, even absent in the sales report.
                if not obj['sales']:
                    for i, label in enumerate(rows[1]):
                        m = month_label(label)
                        if m:
                            obj['sales'][m] = number(row[i])
            else:
                for col in range(3, len(rows[0])):
                    qty = number(row[col])
                    match = re.search(r'поступление до (\d{2}\.\d{2}\.\d{4})', str(rows[0][col]))
                    if qty > 0:
                        eta = datetime.strptime(match.group(1), '%d.%m.%Y').date().isoformat() if match else None
                        obj['transit'].append(dict(qty=qty, eta=eta, source=str(rows[0][col])))
                if 'БУХТАМИ' in str(row[ni]):
                    obj['warnings'].append('Закупка бухтами, учёт в метрах: проверить коэффициент единиц')

        p = bytype['events']
        events = defaultdict(list)
        net = defaultdict(lambda: defaultdict(float))
        for lineno, row in enumerate(read_rows(p), 1):
            if lineno == 1 or not row[3] or not isinstance(row[7], (int, float)):
                continue
            try:
                d = datetime.strptime(str(row[0]), '%d.%m.%Y %H:%M:%S').date()
            except ValueError:
                diagnostics['invalid_event_dates'] += 1
                continue
            if not str(row[2]).startswith('Расходная накладная'):
                diagnostics['non_sales_documents'] += 1
                continue
            if d > date(2026, 9, 22):
                diagnostics['future_events'] += 1
                continue
            obj = item(supplier, row[3], row[4])
            if not obj['unit']:
                obj['unit'] = str(row[5] or '')
            qty = number(row[7])
            diagnostics['transaction_rows'] += 1
            if qty < 0:
                diagnostics['negative_adjustments'] += 1
            net[obj['id']][d.strftime('%Y-%m')] += qty
            if qty > 0:
                events[obj['id']].append(dict(date=d.isoformat(), order_id=str(row[1]), qty=qty))
        for key, evs in events.items():
            obj = catalog[key]
            cleaned = clean_transactions(evs)
            removed = defaultdict(float)
            anomalies = []
            for e in cleaned:
                if e['qty'] > e['clean']:
                    removed[e['date'][:7]] += e['qty']-e['clean']
                    anomalies.append(dict(date=e['date'], qty=e['qty'], regular=e['clean'], reason=e['reason']))
            obj['anomalies'] = sorted(anomalies, key=lambda a: -a['qty'])
            obj['outlier_removed'] = {}
            for m, qty in net[key].items():
                if m < '2025-01':
                    continue  # Earlier event rows are adjustments, not a complete history.
                if m not in obj['sales']:
                    obj['sales'][m] = max(0, qty)
                actual = obj['sales'][m]
                if abs(actual-qty) <= max(1, abs(actual)*.05):
                    obj['outlier_removed'][m] = round(min(actual, removed[m]), 4)
                    diagnostics['reconciled_sku_months'] += 1
                else:
                    diagnostics['unreconciled_sku_months'] += 1
                    if removed[m]:
                        obj['warnings'].append('Есть расхождения динамики и месячных продаж; выбросы в этих месяцах не вычтены')
            obj['sources'].append(f'{p.name} · отбор по коду {obj["code"]}')
            diagnostics['detected_document_outliers'] += len(anomalies)

    result = dict(version=1, as_of='2026-09-22', items=list(catalog.values()), seasonality=seasonality,
                  sources=manifest, diagnostics=dict(diagnostics), limitations=[
                      'В динамике нет ID клиента: доступны выбросы накладных, концентрацию по клиенту проверяем на синтетическом примере.',
                      'Точных периодов stockout нет. Нулевой или пустой месячный остаток не считается доказанным дефицитом.',
                      'Пустые месячные продажи трактуются как отсутствие движения; пустой остаток остаётся неизвестным.',
                      'IEK: остатки на начало сентября, не на дату расчёта. Нужна актуальная ведомость.',
                      'Сроки поставки, цикл заказа и страховой запас заданы как параметры сценария, не как условия поставщиков.',
                      'Сезонность из предоставленных коэффициентов по поставщику, построенных по денежным продажам; нормируется к среднему 1.',
                      'Категории известны для Systeme Electric; правила категорий задаёт менеджер. У IEK справочник категорий отсутствует.',
                      'Коэффициент роста из ведомости SE показан для сверки. Основной тренд рассчитан по очищенным полным месяцам.',
                      'Расчёт по объединённому складу Алматы. Разбивка остатков по складам для обоих поставщиков отсутствует.',
                  ])
    return result


def apply_stockouts(dataset, path):
    lookup = {i['id']: i for i in dataset['items']}
    intervals = defaultdict(set)
    with Path(path).open(encoding='utf-8-sig', newline='') as f:
        for row in csv.DictReader(f):
            key = row['supplier']+':'+row['code']
            if key not in lookup:
                raise ValueError(f'Неизвестный код stockout: {key}')
            start, end = date.fromisoformat(row['start']), date.fromisoformat(row['end'])
            if end < start or (end-start).days > 1000:
                raise ValueError('Неверный интервал stockout')
            while start <= end:
                intervals[key].add(start)
                start += timedelta(days=1)
    for key, days in intervals.items():
        lookup[key]['stockout_days'] = dict(Counter(d.strftime('%Y-%m') for d in days))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, default=Path('data/dataset.json'))
    parser.add_argument('--stockouts', type=Path)
    args = parser.parse_args()
    data = import_workbooks(args.source)
    if args.stockouts:
        apply_stockouts(data, args.stockouts)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(json.dumps(dict(items=len(data['items']), sources=len(data['sources']), diagnostics=data['diagnostics']), ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
